"""
Reportes de solvencia (PROMPT_MODULO_SOLVENCIA.md): catálogo dinámico de
conceptos cobrables, solvencia por grado/mes, estado de pagos por concepto
(resumen agregado + detalle con nombres + Excel) y estado de cuenta del
representante.

Los 4 endpoints principales son de SOLO LECTURA sobre modelos existentes —
no hay migraciones ni escritura. `cobranza/mora.py` sigue siendo la única
fuente de verdad para el criterio de MORA; este módulo no lo duplica, solo
reporta el estado de pago (pagado/parcial/pendiente) de cada concepto.

Definiciones canónicas (no cambiar sin acordarlo antes):
  - Solvente en un mes = tiene la Mensualidad de ese (mes, anio) con
    pagado=True.
  - Denominador ("inscritos") = alumnos con activo=True del grado,
    EXCLUYENDO estatus_financiero='becado'. Mismo criterio de exclusión
    que `ListaMorososView._build_qs` (cobranza/views.py).
"""
import calendar
import re
from datetime import date
from decimal import Decimal

from django.db.models import Count, DecimalField, ExpressionWrapper, F, Q, Sum
from django.http import HttpResponse
from common.busqueda import filtrar_busqueda
from rest_framework import permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView

from config.pagination import StandardResultsPagination

from .models import (
    CuotaInscripcion,
    CuotaProyectoInversion,
    CuotaSolvencia,
    Mensualidad,
    Pago,
    TipoCargoEspecial,
)
from .permissions import filtrar_por_sede
from .serializers import PagoSerializer
from .services import configuracion_activa, meses_ano_escolar, rango_ano_escolar

ROLES_PERMITIDOS = ('director', 'sistemas', 'administrador', 'cobranza', 'cajero')


def _sin_permiso(request):
    """True si el usuario NO tiene uno de los roles permitidos (ni es superuser)."""
    rol = getattr(getattr(request.user, 'perfil', None), 'rol', '')
    return not request.user.is_superuser and rol not in ROLES_PERMITIDOS


# ──────────────────────────────────────────────────────────────────────────────
# BLOQUE 1 — Catálogo dinámico de conceptos cobrables
# ──────────────────────────────────────────────────────────────────────────────

CONCEPTOS_FIJOS = [
    {'clave': 'mensualidad', 'nombre': 'Mensualidad', 'nivel': 'alumno', 'admite_parcial': False, 'periodico': True},
    {'clave': 'inscripcion', 'nombre': 'Inscripción', 'nivel': 'alumno', 'admite_parcial': True, 'periodico': False},
    {'clave': 'solvencia', 'nombre': 'Solvencia', 'nivel': 'alumno', 'admite_parcial': True, 'periodico': False},
]


def conceptos_cobrables():
    """
    Lista completa de conceptos cobrables: los 3 fijos + uno por cada
    TipoCargoEspecial activo (clave `cargo_especial:<id>`). El frontend NO
    decide qué es periódico ni qué admite pago parcial: lo lee de acá.
    """
    conceptos = [dict(c) for c in CONCEPTOS_FIJOS]
    for tipo in TipoCargoEspecial.objects.filter(activo=True).order_by('nombre'):
        conceptos.append({
            'clave': f'cargo_especial:{tipo.id}',
            'nombre': tipo.nombre,
            'nivel': 'representante',
            'admite_parcial': True,
            'periodico': False,
            'tipo_cargo_id': tipo.id,
            'numero_cuotas': tipo.numero_cuotas,
        })
    return conceptos


def info_concepto(clave):
    """Metadatos (nombre/nivel/admite_parcial/periodico/...) de `clave`, o None."""
    for c in conceptos_cobrables():
        if c['clave'] == clave:
            return c
    return None


def resolver_concepto(clave):
    """
    Única función que traduce una clave de concepto a (modelo, nivel,
    filtro_extra). BLOQUE 2, 3 y 4 la usan — no repetir este if/elif en
    varios sitios.

    - modelo: clase del modelo Django a consultar.
    - nivel: 'alumno' o 'representante'.
    - filtro_extra: dict de filtros adicionales para el queryset base del
      modelo (ej. {'tipo_concepto_id': N} para cargos especiales).

    Lanza ValueError si la clave no existe o el TipoCargoEspecial referido
    no está activo.
    """
    if clave == 'mensualidad':
        return Mensualidad, 'alumno', {}
    if clave == 'inscripcion':
        return CuotaInscripcion, 'alumno', {}
    if clave == 'solvencia':
        return CuotaSolvencia, 'alumno', {}
    if clave and clave.startswith('cargo_especial:'):
        try:
            tipo_id = int(clave.split(':', 1)[1])
        except (ValueError, IndexError):
            raise ValueError(f"Clave de concepto inválida: {clave}")
        if not TipoCargoEspecial.objects.filter(id=tipo_id, activo=True).exists():
            raise ValueError(f"Tipo de cargo especial no encontrado o inactivo: {tipo_id}")
        return CuotaProyectoInversion, 'representante', {'tipo_concepto_id': tipo_id}
    raise ValueError(f"Concepto desconocido: {clave}")


class ConceptosCobrablesView(APIView):
    """GET /api/cobranza/conceptos-cobrables/ — catálogo dinámico de conceptos."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        return Response({'conceptos': conceptos_cobrables()})


# ──────────────────────────────────────────────────────────────────────────────
# BLOQUE 2 — Solvencia por grado y mes
# ──────────────────────────────────────────────────────────────────────────────

class SolvenciaMensualView(APIView):
    """
    GET /api/cobranza/solvencia-mensual/

    Reporte AGREGADO y ANÓNIMO. INVARIANTE: la respuesta NO contiene ningún
    nombre propio ni cédula — es el reporte compartible. No agregarle datos
    personales.

    Solvente en un mes = tiene la Mensualidad de ese (mes, anio) con
    pagado=True. Denominador = alumnos activo=True, excluyendo
    estatus_financiero='becado' (mismo criterio que ListaMorososView).

    Devuelve TODOS los meses del período escolar de una sola vez — el
    frontend cambia de mes en cliente, sin volver a llamar.

    Query params: anio_escolar (opcional; ver nota abajo), sede (opcional).

    NOTA sobre `anio_escolar`: se acepta el parámetro pero no selecciona una
    configuración histórica distinta — `ConfiguracionSistema` es un
    singleton (una sola fila activa, ver secretaria/models.py), no existe
    un registro por año escolar contra el cual buscar otro período. El
    reporte siempre usa el período escolar activo. Ver NOTAS_TECNICAS.md.
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if _sin_permiso(request):
            return Response({'error': 'Sin permiso.'}, status=status.HTTP_403_FORBIDDEN)

        config = configuracion_activa()
        rango = rango_ano_escolar(config)
        if not rango:
            return Response(
                {'error': 'No hay período escolar configurado (falta fecha_inicio_ano_escolar/fecha_fin_ano_escolar).'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        meses = meses_ano_escolar(*rango)

        base_qs = Mensualidad.objects.filter(alumno__activo=True).exclude(alumno__estatus_financiero='becado')

        filtro_meses = Q()
        for mes, anio in meses:
            filtro_meses |= Q(mes=mes, anio=anio)
        base_qs = base_qs.filter(filtro_meses)

        base_qs = filtrar_por_sede(request.user, base_qs, campo='alumno__sede')

        sede_param = request.query_params.get('sede')
        if sede_param:
            base_qs = base_qs.filter(alumno__sede_id=sede_param)

        # Rendimiento obligatorio: máximo 2 queries agregadas, nada de
        # iterar alumnos en Python.
        agregados_mes = {
            (row['mes'], row['anio']): row
            for row in base_qs.values('mes', 'anio').annotate(
                total=Count('id'),
                solventes=Count('id', filter=Q(pagado=True)),
            )
        }
        agregados_grado = list(
            base_qs.filter(alumno__grado_seccion__isnull=False)
            .values('mes', 'anio', 'alumno__grado_seccion')
            .annotate(
                total=Count('id'),
                solventes=Count('id', filter=Q(pagado=True)),
            )
        )

        def _desglose(total, solventes):
            pendientes = total - solventes
            porcentaje = round(solventes / total * 100, 1) if total else 0.0
            return pendientes, porcentaje

        meses_resp = []
        total_alumnos_acum = solventes_acum = 0
        for mes, anio in meses:
            row = agregados_mes.get((mes, anio))
            total = row['total'] if row else 0
            solventes = row['solventes'] if row else 0
            pendientes, porcentaje = _desglose(total, solventes)
            meses_resp.append({
                'mes': mes, 'anio': anio,
                'etiqueta': f"{Mensualidad.MESES[mes - 1][1]} {anio}",
                'total_alumnos': total, 'solventes': solventes,
                'pendientes': pendientes, 'porcentaje': porcentaje,
            })
            total_alumnos_acum += total
            solventes_acum += solventes

        por_grado_map = {}
        for row in agregados_grado:
            grado = row['alumno__grado_seccion']
            por_grado_map.setdefault(grado, {})[(row['mes'], row['anio'])] = row

        por_grado_resp = []
        for grado in sorted(por_grado_map.keys()):
            datos = por_grado_map[grado]
            meses_grado = []
            for mes, anio in meses:
                row = datos.get((mes, anio))
                total = row['total'] if row else 0
                solventes = row['solventes'] if row else 0
                pendientes, porcentaje = _desglose(total, solventes)
                meses_grado.append({
                    'mes': mes, 'anio': anio,
                    'total_alumnos': total, 'solventes': solventes,
                    'pendientes': pendientes, 'porcentaje': porcentaje,
                })
            por_grado_resp.append({'grado_seccion': grado, 'meses': meses_grado})

        pendientes_acum = total_alumnos_acum - solventes_acum
        porcentaje_acum = round(solventes_acum / total_alumnos_acum * 100, 1) if total_alumnos_acum else 0.0

        return Response({
            'periodo_escolar': config.periodo_escolar_activo,
            'meses': meses_resp,
            'por_grado': por_grado_resp,
            'totales': {
                'total_alumnos': total_alumnos_acum,
                'solventes': solventes_acum,
                'pendientes': pendientes_acum,
                'porcentaje': porcentaje_acum,
            },
        })


# ──────────────────────────────────────────────────────────────────────────────
# BLOQUE 3 — Estado de pagos por concepto
# ──────────────────────────────────────────────────────────────────────────────

def _filtro_estado(estado, tiene_monto_pagado):
    """Traduce el parámetro `estado` a un Q aplicable a nivel de BD."""
    if estado == 'pagado':
        return Q(pagado=True)
    if estado == 'pendiente':
        if tiene_monto_pagado:
            return Q(pagado=False) & Q(monto_pagado__lte=0)
        return Q(pagado=False)
    if estado == 'parcial':
        if tiene_monto_pagado:
            return Q(pagado=False) & Q(monto_pagado__gt=0)
        return Q(pk__in=[])  # este concepto nunca tiene parciales
    return Q()  # 'todos'


def _estado_monto_saldo(obj, tiene_monto_pagado):
    """(estado, monto_pagado, saldo) de una fila, según las reglas del documento:
    sin monto_pagado -> pagado/pendiente; con monto_pagado -> pagado/parcial/pendiente."""
    monto_usd = obj.monto_usd
    monto_pagado = obj.monto_pagado if tiene_monto_pagado else (monto_usd if obj.pagado else Decimal('0.00'))
    saldo = monto_usd - monto_pagado
    if not tiene_monto_pagado:
        estado = 'pagado' if obj.pagado else 'pendiente'
    elif obj.pagado:
        estado = 'pagado'
    elif monto_pagado > 0:
        estado = 'parcial'
    else:
        estado = 'pendiente'
    return estado, monto_pagado, saldo


def _dias_atraso_fila(obj, clave, tiene_monto_pagado, hoy):
    """dias_atraso solo aplica a mensualidad pendiente; None para el resto.
    Misma fórmula que cobranza/mora.py::calcular_dias_atraso (~línea 260)."""
    if clave != 'mensualidad':
        return None
    estado, _, _ = _estado_monto_saldo(obj, tiene_monto_pagado)
    if estado != 'pendiente':
        return None
    alumno = obj.alumno
    dia = min(alumno.dia_limite_pago or 1, calendar.monthrange(obj.anio, obj.mes)[1])
    vencimiento = date(obj.anio, obj.mes, dia)
    return max((hoy - vencimiento).days, 0)


def _agregar_conteos(qs, group_fields, tiene_monto_pagado):
    """
    Agrega `qs` por `group_fields` (o sin agrupar si está vacío) en UNA
    consulta, devolviendo total/pagados/parciales/pendientes/porcentaje/
    monto_pendiente por grupo.
    """
    annotate_kwargs = {
        'total': Count('id'),
        'pagados': Count('id', filter=Q(pagado=True)),
    }
    if tiene_monto_pagado:
        annotate_kwargs['parciales'] = Count('id', filter=Q(pagado=False, monto_pagado__gt=0))
        saldo_expr = ExpressionWrapper(
            F('monto_usd') - F('monto_pagado'), output_field=DecimalField(max_digits=10, decimal_places=2)
        )
    else:
        saldo_expr = F('monto_usd')
    annotate_kwargs['monto_pendiente'] = Sum(saldo_expr, filter=Q(pagado=False))

    if group_fields:
        rows = list(qs.values(*group_fields).annotate(**annotate_kwargs))
    else:
        rows = [qs.aggregate(**annotate_kwargs)]

    for r in rows:
        if not tiene_monto_pagado:
            r['parciales'] = 0
        r['pendientes'] = r['total'] - r['pagados'] - r['parciales']
        r['porcentaje'] = round(r['pagados'] / r['total'] * 100, 1) if r['total'] else 0.0
        r['monto_pendiente'] = r['monto_pendiente'] or Decimal('0.00')
    return rows


def _resolver_y_filtrar(request):
    """
    Punto único de resolución + filtrado compartido por ResumenPorConceptoView,
    EstadoPorConceptoView y ExportarEstadoPorConceptoExcelView (todo menos el
    filtro `estado`, que cada vista aplica por separado según lo necesite).

    Devuelve (qs, nivel, tiene_monto_pagado, tiene_numero_cuota, info, clave, None)
    o (None, None, None, None, None, None, Response) si hay un error de validación.
    """
    clave = request.query_params.get('concepto')
    if not clave:
        return None, None, None, None, None, None, Response(
            {'error': 'El parámetro concepto es obligatorio.'}, status=status.HTTP_400_BAD_REQUEST)
    info = info_concepto(clave)
    if not info:
        return None, None, None, None, None, None, Response(
            {'error': f'Concepto desconocido: {clave}'}, status=status.HTTP_400_BAD_REQUEST)
    try:
        modelo, nivel, filtro_extra = resolver_concepto(clave)
    except ValueError as exc:
        return None, None, None, None, None, None, Response({'error': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    tiene_monto_pagado = hasattr(modelo, 'monto_pagado')
    tiene_numero_cuota = hasattr(modelo, 'numero_cuota')

    select_related = ('alumno', 'alumno__representante') if nivel == 'alumno' else ('representante',)
    qs = modelo.objects.filter(**filtro_extra).select_related(*select_related)
    if nivel == 'representante':
        qs = qs.prefetch_related('representante__alumnos')

    campo_sede = 'alumno__sede' if nivel == 'alumno' else 'representante__alumnos__sede'
    qs = filtrar_por_sede(request.user, qs, campo=campo_sede)

    sede_param = request.query_params.get('sede')
    if sede_param:
        campo = 'alumno__sede_id' if nivel == 'alumno' else 'representante__alumnos__sede_id'
        qs = qs.filter(**{campo: sede_param}).distinct()

    if hasattr(modelo, 'periodo_escolar'):
        periodo_escolar_param = request.query_params.get('periodo_escolar')
        if periodo_escolar_param:
            qs = qs.filter(periodo_escolar=periodo_escolar_param)
        else:
            config = configuracion_activa()
            if config:
                qs = qs.filter(periodo_escolar=config.periodo_escolar_activo)

    if info['periodico']:
        mes_param = request.query_params.get('mes')
        anio_param = request.query_params.get('anio')
        if mes_param:
            qs = qs.filter(mes=mes_param)
        if anio_param:
            qs = qs.filter(anio=anio_param)

    if tiene_numero_cuota:
        numero_cuota_param = request.query_params.get('numero_cuota')
        if numero_cuota_param:
            qs = qs.filter(numero_cuota=numero_cuota_param)

    grado_seccion_param = request.query_params.get('grado_seccion')
    if grado_seccion_param:
        campo = 'alumno__grado_seccion' if nivel == 'alumno' else 'representante__alumnos__grado_seccion'
        qs = qs.filter(**{campo: grado_seccion_param}).distinct()

    buscar = (request.query_params.get('buscar') or '').strip()
    if buscar:
        if nivel == 'alumno':
            qs = filtrar_busqueda(qs, buscar, [
                'alumno__nombre', 'alumno__apellido', 'alumno__cedula_escolar',
                'alumno__representante__nombre', 'alumno__representante__apellido',
                'alumno__representante__cedula',
            ]).distinct()
        else:
            qs = filtrar_busqueda(qs, buscar, [
                'representante__nombre', 'representante__apellido', 'representante__cedula',
                'representante__alumnos__nombre', 'representante__alumnos__apellido',
            ]).distinct()

    return qs, nivel, tiene_monto_pagado, tiene_numero_cuota, info, clave, None


class ResumenPorConceptoView(APIView):
    """
    GET /api/cobranza/estado-por-concepto/resumen/

    Resumen agregado (sin nombres) de un concepto: desglose por mes si es
    periódico, o una sola línea (una por cuota, si el concepto tiene varias)
    si no lo es. El filtrado de "no mostrar lo ya cobrado al 100%" lo hace
    el frontend — este endpoint siempre devuelve todo.
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if _sin_permiso(request):
            return Response({'error': 'Sin permiso.'}, status=status.HTTP_403_FORBIDDEN)

        vista = request.query_params.get('vista', 'global')
        if vista not in ('global', 'grado'):
            return Response({'error': "vista debe ser 'global' o 'grado'."}, status=status.HTTP_400_BAD_REQUEST)

        qs, nivel, tiene_monto_pagado, tiene_numero_cuota, info, clave, error = _resolver_y_filtrar(request)
        if error:
            return error

        campo_grado = 'alumno__grado_seccion' if nivel == 'alumno' else 'representante__alumnos__grado_seccion'

        def _con_grados(filas):
            return sorted((
                {
                    'grado_seccion': f[campo_grado],
                    'total': f['total'], 'pagados': f['pagados'],
                    'pendientes': f['total'] - f['pagados'] - f['parciales'],
                    'porcentaje': round(f['pagados'] / f['total'] * 100, 1) if f['total'] else 0.0,
                }
                for f in filas if f[campo_grado]
            ), key=lambda g: g['grado_seccion'])

        lineas = []
        total_acc = pagados_acc = parciales_acc = pendientes_acc = 0

        if info['periodico']:
            config = configuracion_activa()
            rango = rango_ano_escolar(config)
            if not rango:
                return Response({'error': 'No hay período escolar configurado.'}, status=status.HTTP_400_BAD_REQUEST)
            meses = meses_ano_escolar(*rango)
            filtro_meses = Q()
            for mes, anio in meses:
                filtro_meses |= Q(mes=mes, anio=anio)
            qs = qs.filter(filtro_meses)

            group_fields = ['mes', 'anio'] + ([campo_grado] if vista == 'grado' else [])
            filas = _agregar_conteos(qs, group_fields, tiene_monto_pagado)
            por_mes = {}
            for f in filas:
                por_mes.setdefault((f['mes'], f['anio']), []).append(f)

            for mes, anio in meses:
                filas_mes = por_mes.get((mes, anio), [])
                total = sum(f['total'] for f in filas_mes)
                pagados = sum(f['pagados'] for f in filas_mes)
                parciales = sum(f['parciales'] for f in filas_mes)
                pendientes = total - pagados - parciales
                porcentaje = round(pagados / total * 100, 1) if total else 0.0
                monto_pendiente = sum((f['monto_pendiente'] for f in filas_mes), Decimal('0.00'))
                linea = {
                    'etiqueta': f"{Mensualidad.MESES[mes - 1][1]} {anio}", 'mes': mes, 'anio': anio,
                    'total': total, 'pagados': pagados, 'parciales': parciales,
                    'pendientes': pendientes, 'porcentaje': porcentaje,
                    'monto_pendiente_usd': str(monto_pendiente),
                }
                if vista == 'grado':
                    linea['grados'] = _con_grados(filas_mes)
                lineas.append(linea)
                total_acc += total; pagados_acc += pagados
                parciales_acc += parciales; pendientes_acc += pendientes
        else:
            group_fields = (['numero_cuota'] if tiene_numero_cuota else []) + ([campo_grado] if vista == 'grado' else [])
            filas = _agregar_conteos(qs, group_fields, tiene_monto_pagado)

            if tiene_numero_cuota:
                por_cuota = {}
                for f in filas:
                    por_cuota.setdefault(f['numero_cuota'], []).append(f)
                claves_cuota = sorted(por_cuota.keys())
            else:
                por_cuota = {None: filas}
                claves_cuota = [None]

            for numero_cuota in claves_cuota:
                filas_cuota = por_cuota[numero_cuota]
                total = sum(f['total'] for f in filas_cuota)
                pagados = sum(f['pagados'] for f in filas_cuota)
                parciales = sum(f['parciales'] for f in filas_cuota)
                pendientes = total - pagados - parciales
                porcentaje = round(pagados / total * 100, 1) if total else 0.0
                monto_pendiente = sum((f['monto_pendiente'] for f in filas_cuota), Decimal('0.00'))
                etiqueta = info['nombre']
                if numero_cuota is not None:
                    etiqueta = f"{info['nombre']} — Cuota {numero_cuota} de {info.get('numero_cuotas', numero_cuota)}"
                linea = {
                    'etiqueta': etiqueta, 'mes': None, 'anio': None,
                    'total': total, 'pagados': pagados, 'parciales': parciales,
                    'pendientes': pendientes, 'porcentaje': porcentaje,
                    'monto_pendiente_usd': str(monto_pendiente),
                }
                if vista == 'grado':
                    linea['grados'] = _con_grados(filas_cuota)
                lineas.append(linea)
                total_acc += total; pagados_acc += pagados
                parciales_acc += parciales; pendientes_acc += pendientes

        return Response({
            'concepto': clave, 'concepto_nombre': info['nombre'],
            'periodico': info['periodico'], 'nivel': nivel, 'vista': vista,
            'lineas': lineas,
            'totales': {
                'total': total_acc, 'pagados': pagados_acc, 'parciales': parciales_acc,
                'pendientes': pendientes_acc,
                'porcentaje': round(pagados_acc / total_acc * 100, 1) if total_acc else 0.0,
            },
        })


def _mensualidades_por_monto_qs(request):
    """Base única para el reporte de montos cobrados por mensualidad.

    El reporte representa el cargo registrado (``monto_usd``), no el importe
    de un pago. Así, un alumno con descuento conserva el monto que realmente
    se le cobra y puede clasificarse de forma independiente.
    """
    mes = request.query_params.get('mes')
    anio = request.query_params.get('anio')
    if not mes or not anio:
        return None, Response(
            {'error': 'Los parámetros mes y anio son obligatorios.'},
            status=status.HTTP_400_BAD_REQUEST,
        )
    try:
        mes_num, anio_num = int(mes), int(anio)
    except (TypeError, ValueError):
        return None, Response({'error': 'mes y anio deben ser números válidos.'}, status=status.HTTP_400_BAD_REQUEST)
    if not 1 <= mes_num <= 12:
        return None, Response({'error': 'mes debe estar entre 1 y 12.'}, status=status.HTTP_400_BAD_REQUEST)

    # `alumno__activo=True`: los retirados/eliminados conservan sus Mensualidad
    # (soft delete) y, al reinscribir a un alumno, aparecen como filas extra
    # del mismo nombre con otro monto — es lo que duplicaba listados y Excel/PDF.
    qs = Mensualidad.objects.filter(
        mes=mes_num, anio=anio_num, monto_usd__gt=0, alumno__activo=True,
    ).select_related(
        'alumno', 'alumno__representante'
    )
    qs = filtrar_por_sede(request.user, qs, campo='alumno__sede')
    if sede := request.query_params.get('sede'):
        qs = qs.filter(alumno__sede_id=sede)

    grado = request.query_params.get('grado_seccion')
    if grado:
        qs = qs.filter(alumno__grado_seccion=grado)
    buscar = (request.query_params.get('buscar') or '').strip()
    if buscar:
        qs = filtrar_busqueda(qs, buscar, [
            'alumno__nombre', 'alumno__apellido', 'alumno__cedula_escolar',
            'alumno__representante__nombre', 'alumno__representante__apellido',
            'alumno__representante__cedula',
        ]).distinct()
    estado = request.query_params.get('estado', 'todos')
    if estado not in ('todos', 'pagado', 'pendiente'):
        return None, Response({'error': 'estado debe ser todos, pagado o pendiente.'}, status=status.HTTP_400_BAD_REQUEST)
    if estado != 'todos':
        qs = qs.filter(pagado=(estado == 'pagado'))
    return qs, None


def _serializar_mensualidad_por_monto(mensualidad):
    alumno = mensualidad.alumno
    representante = alumno.representante
    pagado = mensualidad.monto_usd if mensualidad.pagado else Decimal('0.00')
    return {
        'alumno_id': alumno.id,
        'nombre': f'{alumno.nombre} {alumno.apellido}',
        'cedula_escolar': alumno.cedula_escolar,
        'grado_seccion': alumno.grado_seccion,
        'representante': {
            'id': representante.id,
            'nombre': f'{representante.nombre} {representante.apellido}',
            'cedula': representante.cedula,
            'telefono': representante.telefono,
        } if representante else None,
        'monto_usd': str(mensualidad.monto_usd),
        'monto_pagado_usd': str(pagado),
        'saldo_usd': str(mensualidad.monto_usd - pagado),
        'estado': 'pagado' if mensualidad.pagado else 'pendiente',
        'fecha_pago': mensualidad.fecha_pago,
    }


class MensualidadesPorMontoView(APIView):
    """Resumen por monto de mensualidad de un período, sin datos personales."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if _sin_permiso(request):
            return Response({'error': 'Sin permiso.'}, status=status.HTTP_403_FORBIDDEN)
        qs, error = _mensualidades_por_monto_qs(request)
        if error:
            return error
        grupos = qs.values('monto_usd').annotate(
            estudiantes=Count('alumno_id', distinct=True), total_facturado_usd=Sum('monto_usd'),
        ).order_by('monto_usd')
        return Response({
            'mes': int(request.query_params['mes']), 'anio': int(request.query_params['anio']),
            'grupos': [{
                'monto_usd': str(g['monto_usd']), 'estudiantes': g['estudiantes'],
                'total_facturado_usd': str(g['total_facturado_usd']),
            } for g in grupos],
        })


class DetalleMensualidadesPorMontoView(APIView):
    """Alumnos de un importe seleccionado, con filtros y paginación."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if _sin_permiso(request):
            return Response({'error': 'Sin permiso.'}, status=status.HTTP_403_FORBIDDEN)
        qs, error = _mensualidades_por_monto_qs(request)
        if error:
            return error
        try:
            monto = Decimal(request.query_params.get('monto', ''))
        except Exception:
            return Response({'error': 'monto es obligatorio y debe ser válido.'}, status=status.HTTP_400_BAD_REQUEST)
        qs = qs.filter(monto_usd=monto).order_by('alumno__apellido', 'alumno__nombre', 'id')
        resumen = qs.aggregate(
            estudiantes=Count('alumno_id', distinct=True), total_facturado_usd=Sum('monto_usd'),
            pagados=Count('id', filter=Q(pagado=True)), pendientes=Count('id', filter=Q(pagado=False)),
        )
        paginator = StandardResultsPagination()
        pagina = paginator.paginate_queryset(qs, request, view=self)
        response = paginator.get_paginated_response([_serializar_mensualidad_por_monto(item) for item in pagina])
        response.data['resumen'] = {
            'monto_usd': str(monto), 'estudiantes': resumen['estudiantes'],
            'total_facturado_usd': str(resumen['total_facturado_usd'] or Decimal('0.00')),
            'pagados': resumen['pagados'], 'pendientes': resumen['pendientes'],
        }
        return response


def _clave_orden_grado(grado_seccion):
    """Orden académico natural: Inicial (Sala/Maternal/Preescolar) → Grado → Año.

    No existe un orden de grados persistido (ConfiguracionGrado no tiene
    campo de orden y el resto del backend ordena alfabéticamente), así que se
    deduce del texto: nivel, luego el primer número ("1er", "2do"…) y por
    último el texto completo para desempatar secciones ("1er Grado A" < "B").
    """
    texto = grado_seccion.lower()
    if any(p in texto for p in ('sala', 'maternal', 'preescolar', 'inicial')):
        nivel = 0
    elif 'grado' in texto:
        nivel = 1
    elif 'año' in texto or 'ano' in texto:
        nivel = 2
    else:
        nivel = 3
    numero = re.search(r'\d+', texto)
    return (nivel, int(numero.group()) if numero else 0, texto)


class MensualidadesPorGradoView(APIView):
    """Mensualidades de un período agrupadas por grado, con nombres (modal "Ver detalles").

    Sin paginación: un período cabe completo en memoria. Una sola query
    (select_related en la base) y el agrupado se hace en Python. Los alumnos
    sin grado van al final con etiqueta vacía.
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if _sin_permiso(request):
            return Response({'error': 'Sin permiso.'}, status=status.HTTP_403_FORBIDDEN)
        qs, error = _mensualidades_por_monto_qs(request)
        if error:
            return error
        monto_param = request.query_params.get('monto')
        if monto_param:
            try:
                qs = qs.filter(monto_usd=Decimal(monto_param))
            except Exception:
                return Response({'error': 'monto debe ser válido.'}, status=status.HTTP_400_BAD_REQUEST)

        grupos = {}
        for item in qs.order_by('alumno__apellido', 'alumno__nombre', 'id'):
            alumno, representante = item.alumno, item.alumno.representante
            grupo = grupos.setdefault(alumno.grado_seccion or '', {'ids': set(), 'total': Decimal('0.00'), 'alumnos': []})
            grupo['ids'].add(alumno.id)
            grupo['total'] += item.monto_usd
            grupo['alumnos'].append({
                'alumno_id': alumno.id, 'nombre': alumno.nombre, 'apellido': alumno.apellido,
                'representante': f'{representante.nombre} {representante.apellido}' if representante else None,
                'monto_usd': str(item.monto_usd), 'estado': 'pagado' if item.pagado else 'pendiente',
            })

        # Los sin grado ('') al final; el resto en orden académico natural.
        orden = sorted((g for g in grupos if g), key=_clave_orden_grado) + ([''] if '' in grupos else [])
        estudiantes = set().union(*(g['ids'] for g in grupos.values()))
        return Response({
            'mes': int(request.query_params['mes']), 'anio': int(request.query_params['anio']),
            'resumen': {
                'estudiantes': len(estudiantes),
                'total_usd': str(sum((g['total'] for g in grupos.values()), Decimal('0.00'))),
                'grados': len(grupos),
            },
            'grados': [{
                'grado_seccion': grado, 'estudiantes': len(grupos[grado]['ids']),
                'total_usd': str(grupos[grado]['total']), 'alumnos': grupos[grado]['alumnos'],
            } for grado in orden],
        })


class ExportarMensualidadesPorMontoExcelView(APIView):
    """Excel legible: resumen y detalle con anchos, filtros y panel congelado."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if _sin_permiso(request):
            return Response({'error': 'Sin permiso.'}, status=status.HTTP_403_FORBIDDEN)
        qs, error = _mensualidades_por_monto_qs(request)
        if error:
            return error
        monto_param = request.query_params.get('monto')
        if monto_param:
            try:
                qs = qs.filter(monto_usd=Decimal(monto_param))
            except Exception:
                return Response({'error': 'monto debe ser válido.'}, status=status.HTTP_400_BAD_REQUEST)

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        libro = Workbook()
        resumen_ws = libro.active
        resumen_ws.title = 'Resumen por monto'
        detalle_ws = libro.create_sheet('Detalle de alumnos')
        encabezado = PatternFill('solid', fgColor='1F4E79')
        fuente_encabezado = Font(color='FFFFFF', bold=True)

        grupos = list(qs.values('monto_usd').annotate(
            estudiantes=Count('alumno_id', distinct=True), total_facturado_usd=Sum('monto_usd'),
        ).order_by('monto_usd'))
        resumen_ws.append(['Mensualidad (USD)', 'Estudiantes', 'Total facturado (USD)'])
        for grupo in grupos:
            resumen_ws.append([float(grupo['monto_usd']), grupo['estudiantes'], float(grupo['total_facturado_usd'])])
        detalle_ws.append(['Nombre', 'Apellido', 'Cédula escolar', 'Grado / Sección', 'Representante', 'Teléfono', 'Mensualidad (USD)', 'Pagado (USD)', 'Saldo (USD)', 'Estado', 'Fecha de pago'])
        for item in qs.order_by('monto_usd', 'alumno__apellido', 'alumno__nombre'):
            alumno, representante = item.alumno, item.alumno.representante
            pagado = item.monto_usd if item.pagado else Decimal('0.00')
            detalle_ws.append([
                alumno.nombre, alumno.apellido, alumno.cedula_escolar, alumno.grado_seccion,
                f'{representante.nombre} {representante.apellido}' if representante else '',
                representante.telefono if representante else '', float(item.monto_usd), float(pagado),
                float(item.monto_usd - pagado), 'Pagado' if item.pagado else 'Pendiente', item.fecha_pago or '',
            ])
        for hoja in (resumen_ws, detalle_ws):
            hoja.freeze_panes = 'A2'
            hoja.auto_filter.ref = hoja.dimensions
            for celda in hoja[1]:
                celda.fill = encabezado
                celda.font = fuente_encabezado
                celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            for columna in range(1, hoja.max_column + 1):
                largo = max((len(str(hoja.cell(fila, columna).value or '')) for fila in range(1, hoja.max_row + 1)), default=10)
                hoja.column_dimensions[get_column_letter(columna)].width = min(max(largo + 2, 12), 32)
            for fila in hoja.iter_rows(min_row=2):
                for celda in fila:
                    celda.alignment = Alignment(vertical='top', wrap_text=True)
            for columna in (1, 3) if hoja is resumen_ws else (7, 8, 9):
                for celda in hoja.iter_cols(min_col=columna, max_col=columna, min_row=2):
                    for valor in celda:
                        valor.number_format = '$#,##0.00'
        respuesta = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        respuesta['Content-Disposition'] = f'attachment; filename="mensualidades_por_monto_{request.query_params["anio"]}_{request.query_params["mes"]}.xlsx"'
        libro.save(respuesta)
        return respuesta


class EstadoPorConceptoView(APIView):
    """
    GET /api/cobranza/estado-por-concepto/

    Detalle CON nombres de quién pagó / quién debe un concepto, paginado
    (StandardResultsPagination). `resumen` se calcula sobre el queryset
    completo con los mismos filtros salvo `estado` (para poder mostrar el
    desglose pagado/parcial/pendiente completo aunque la página actual esté
    filtrada a un solo estado).
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if _sin_permiso(request):
            return Response({'error': 'Sin permiso.'}, status=status.HTTP_403_FORBIDDEN)

        estado_param = request.query_params.get('estado', 'todos')
        if estado_param not in ('pagado', 'pendiente', 'parcial', 'todos'):
            return Response({'error': "estado debe ser pagado, pendiente, parcial o todos."}, status=status.HTTP_400_BAD_REQUEST)

        qs, nivel, tiene_monto_pagado, tiene_numero_cuota, info, clave, error = _resolver_y_filtrar(request)
        if error:
            return error

        resumen = self._calcular_resumen(qs, tiene_monto_pagado)

        qs = qs.filter(_filtro_estado(estado_param, tiene_monto_pagado)).order_by('id')

        paginator = StandardResultsPagination()
        pagina = paginator.paginate_queryset(qs, request, view=self)

        hoy = date.today()
        resultados = [self._serializar_fila(obj, nivel, clave, tiene_monto_pagado, hoy) for obj in pagina]
        response = paginator.get_paginated_response(resultados)
        response.data['resumen'] = resumen
        return response

    @staticmethod
    def _calcular_resumen(qs, tiene_monto_pagado):
        total = qs.count()
        pagados = qs.filter(pagado=True).count()
        if tiene_monto_pagado:
            parciales = qs.filter(pagado=False, monto_pagado__gt=0).count()
            saldo_expr = ExpressionWrapper(
                F('monto_usd') - F('monto_pagado'), output_field=DecimalField(max_digits=10, decimal_places=2)
            )
            cobrado = qs.aggregate(t=Sum('monto_pagado'))['t'] or Decimal('0.00')
        else:
            parciales = 0
            saldo_expr = F('monto_usd')
            cobrado = qs.filter(pagado=True).aggregate(t=Sum('monto_usd'))['t'] or Decimal('0.00')
        pendiente = qs.filter(pagado=False).aggregate(t=Sum(saldo_expr))['t'] or Decimal('0.00')
        pendientes = total - pagados - parciales
        return {
            'total_filas': total, 'pagados': pagados, 'parciales': parciales, 'pendientes': pendientes,
            'monto_cobrado_usd': str(cobrado), 'monto_pendiente_usd': str(pendiente),
        }

    @staticmethod
    def _serializar_fila(obj, nivel, clave, tiene_monto_pagado, hoy):
        estado, monto_pagado, saldo = _estado_monto_saldo(obj, tiene_monto_pagado)
        dias_atraso = _dias_atraso_fila(obj, clave, tiene_monto_pagado, hoy)

        if nivel == 'alumno':
            alumno = obj.alumno
            representante = alumno.representante
            return {
                'nivel': 'alumno',
                'alumno_id': alumno.id,
                'nombre': f"{alumno.nombre} {alumno.apellido}",
                'cedula_escolar': alumno.cedula_escolar,
                'grado_seccion': alumno.grado_seccion,
                'representante': {
                    'id': representante.id,
                    'nombre': f"{representante.nombre} {representante.apellido}",
                    'cedula': representante.cedula,
                    'telefono': representante.telefono,
                } if representante else None,
                'monto_usd': str(obj.monto_usd),
                'monto_pagado_usd': str(monto_pagado),
                'saldo_usd': str(saldo),
                'estado': estado,
                'fecha_pago': obj.fecha_pago,
                'dias_atraso': dias_atraso,
            }

        representante = obj.representante
        return {
            'nivel': 'representante',
            'representante_id': representante.id,
            'nombre': f"{representante.nombre} {representante.apellido}",
            'cedula': representante.cedula,
            'telefono': representante.telefono,
            'alumnos': [f"{a.nombre} {a.apellido}" for a in representante.alumnos.all()],
            'numero_cuota': getattr(obj, 'numero_cuota', None),
            'monto_usd': str(obj.monto_usd),
            'monto_pagado_usd': str(monto_pagado),
            'saldo_usd': str(saldo),
            'estado': estado,
            'fecha_pago': obj.fecha_pago,
            'dias_atraso': None,
        }


class ExportarEstadoPorConceptoExcelView(APIView):
    """
    GET /api/cobranza/estado-por-concepto/exportar-excel/

    Mismos filtros que EstadoPorConceptoView, sin paginar, vía ExcelExporter.
    Sigue el patrón exacto de ExportarMorososExcelView (cobranza/views.py
    ~línea 3061): solo IsAuthenticated, sin chequeo de rol adicional.
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        from cobranza.exports import ExcelExporter

        estado_param = request.query_params.get('estado', 'todos')
        if estado_param not in ('pagado', 'pendiente', 'parcial', 'todos'):
            return Response({'error': "estado debe ser pagado, pendiente, parcial o todos."}, status=status.HTTP_400_BAD_REQUEST)

        qs, nivel, tiene_monto_pagado, tiene_numero_cuota, info, clave, error = _resolver_y_filtrar(request)
        if error:
            return error
        qs = qs.filter(_filtro_estado(estado_param, tiene_monto_pagado)).order_by('id')

        hoy = date.today()

        if nivel == 'alumno':
            columns = [
                ('Nombre', lambda o: o.alumno.nombre),
                ('Apellido', lambda o: o.alumno.apellido),
                ('Cédula Escolar', lambda o: o.alumno.cedula_escolar),
                ('Grado / Sección', lambda o: o.alumno.grado_seccion),
                ('Representante', lambda o: f"{o.alumno.representante.nombre} {o.alumno.representante.apellido}" if o.alumno.representante_id else ''),
                ('Tel. Representante', lambda o: o.alumno.representante.telefono if o.alumno.representante_id else ''),
                ('Monto (USD)', lambda o: o.monto_usd),
                ('Monto Pagado (USD)', lambda o: _estado_monto_saldo(o, tiene_monto_pagado)[1]),
                ('Saldo (USD)', lambda o: _estado_monto_saldo(o, tiene_monto_pagado)[2]),
                ('Estado', lambda o: _estado_monto_saldo(o, tiene_monto_pagado)[0]),
                ('Fecha de Pago', lambda o: o.fecha_pago),
                ('Días de Atraso', lambda o: _dias_atraso_fila(o, clave, tiene_monto_pagado, hoy)),
            ]
        else:
            columns = [
                ('Representante', lambda o: f"{o.representante.nombre} {o.representante.apellido}"),
                ('Cédula', lambda o: o.representante.cedula),
                ('Teléfono', lambda o: o.representante.telefono),
                ('Alumnos', lambda o: ', '.join(f"{a.nombre} {a.apellido}" for a in o.representante.alumnos.all())),
                ('N° Cuota', lambda o: getattr(o, 'numero_cuota', None)),
                ('Monto (USD)', lambda o: o.monto_usd),
                ('Monto Pagado (USD)', lambda o: _estado_monto_saldo(o, tiene_monto_pagado)[1]),
                ('Saldo (USD)', lambda o: _estado_monto_saldo(o, tiene_monto_pagado)[2]),
                ('Estado', lambda o: _estado_monto_saldo(o, tiene_monto_pagado)[0]),
                ('Fecha de Pago', lambda o: o.fecha_pago),
            ]

        return ExcelExporter.export(qs, columns, f"estado_{clave.replace(':', '_')}_{hoy}")


# ──────────────────────────────────────────────────────────────────────────────
# BLOQUE 4 — Estado de cuenta del representante
# ──────────────────────────────────────────────────────────────────────────────

class EstadoCuentaRepresentanteView(APIView):
    """
    GET /api/cobranza/representantes/<id>/estado-cuenta/

    `cargos` recorre TODOS los conceptos del catálogo (BLOQUE 1):
    mensualidades/inscripción/solvencia de los alumnos activos del
    representante, más sus cargos especiales. `historial_pagos` (paginado)
    incluye pagos anulados (no se descuentan de pagado_total_usd: como
    `anular_pago` revierte pagado=False en las cuotas/mensualidades
    afectadas, ver cobranza/correcciones.py, el total calculado a partir de
    esas cuotas ya excluye lo anulado sin necesidad de restarlo aparte).
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, representante_id):
        if _sin_permiso(request):
            return Response({'error': 'Sin permiso.'}, status=status.HTTP_403_FORBIDDEN)

        from secretaria.models import Representante

        representante_qs = filtrar_por_sede(
            request.user, Representante.objects.all(), campo='alumnos__sede'
        ).distinct()
        representante = representante_qs.filter(pk=representante_id).first()
        if not representante:
            return Response({'error': 'Representante no encontrado.'}, status=status.HTTP_404_NOT_FOUND)

        alumnos = list(representante.alumnos.all().order_by('nombre', 'apellido'))
        alumnos_data = [
            {'id': a.id, 'nombre': f"{a.nombre} {a.apellido}", 'grado_seccion': a.grado_seccion, 'activo': a.activo}
            for a in alumnos
        ]

        cargos = []
        deuda_total = Decimal('0.00')
        pagado_total = Decimal('0.00')
        cargos_pendientes = 0

        # Mensualidades — solo de alumnos activos (ver docstring de la clase).
        mensualidades = list(
            Mensualidad.objects.filter(alumno__representante=representante, alumno__activo=True)
            .select_related('alumno').order_by('anio', 'mes')
        )
        if mensualidades:
            items, subtotal, subtotal_pagado, pendientes = [], Decimal('0.00'), Decimal('0.00'), 0
            for m in mensualidades:
                monto_pagado = m.monto_usd if m.pagado else Decimal('0.00')
                saldo = m.monto_usd - monto_pagado
                items.append({
                    'descripcion': f"{Mensualidad.MESES[m.mes - 1][1]} {m.anio}",
                    'alumno': f"{m.alumno.nombre} {m.alumno.apellido}",
                    'monto_usd': str(m.monto_usd), 'monto_pagado_usd': str(monto_pagado),
                    'saldo_usd': str(saldo), 'estado': 'pagado' if m.pagado else 'pendiente',
                    'fecha_pago': m.fecha_pago,
                })
                subtotal += m.monto_usd
                subtotal_pagado += monto_pagado
                pendientes += 0 if m.pagado else 1
            cargos.append({
                'concepto': 'mensualidad', 'concepto_nombre': 'Mensualidad', 'nivel': 'alumno',
                'items': items, 'subtotal_usd': str(subtotal), 'subtotal_pagado_usd': str(subtotal_pagado),
                'saldo_usd': str(subtotal - subtotal_pagado), 'pendientes': pendientes,
            })
            deuda_total += subtotal - subtotal_pagado
            pagado_total += subtotal_pagado
            cargos_pendientes += pendientes

        # Inscripción — solo de alumnos activos.
        inscripciones = list(
            CuotaInscripcion.objects.filter(alumno__representante=representante, alumno__activo=True)
            .select_related('alumno').order_by('-periodo_escolar')
        )
        if inscripciones:
            items, subtotal, subtotal_pagado, pendientes = [], Decimal('0.00'), Decimal('0.00'), 0
            for c in inscripciones:
                monto_pagado = c.monto_pagado
                saldo = c.monto_usd - monto_pagado
                items.append({
                    'descripcion': f"Inscripción {c.periodo_escolar}",
                    'alumno': f"{c.alumno.nombre} {c.alumno.apellido}",
                    'monto_usd': str(c.monto_usd), 'monto_pagado_usd': str(monto_pagado),
                    'saldo_usd': str(saldo), 'estado': 'pagado' if c.pagado else 'pendiente',
                    'fecha_pago': c.fecha_pago,
                })
                subtotal += c.monto_usd
                subtotal_pagado += monto_pagado
                pendientes += 0 if c.pagado else 1
            cargos.append({
                'concepto': 'inscripcion', 'concepto_nombre': 'Inscripción', 'nivel': 'alumno',
                'items': items, 'subtotal_usd': str(subtotal), 'subtotal_pagado_usd': str(subtotal_pagado),
                'saldo_usd': str(subtotal - subtotal_pagado), 'pendientes': pendientes,
            })
            deuda_total += subtotal - subtotal_pagado
            pagado_total += subtotal_pagado
            cargos_pendientes += pendientes

        # Solvencia — solo de alumnos activos.
        solvencias = list(
            CuotaSolvencia.objects.filter(alumno__representante=representante, alumno__activo=True)
            .select_related('alumno').order_by('-periodo_escolar')
        )
        if solvencias:
            items, subtotal, subtotal_pagado, pendientes = [], Decimal('0.00'), Decimal('0.00'), 0
            for s in solvencias:
                saldo = s.monto_usd - s.monto_pagado
                estado = 'pagado' if s.pagado else ('parcial' if s.monto_pagado > 0 else 'pendiente')
                items.append({
                    'descripcion': f"Solvencia {s.periodo_escolar}",
                    'alumno': f"{s.alumno.nombre} {s.alumno.apellido}",
                    'monto_usd': str(s.monto_usd), 'monto_pagado_usd': str(s.monto_pagado),
                    'saldo_usd': str(saldo), 'estado': estado, 'fecha_pago': s.fecha_pago,
                })
                subtotal += s.monto_usd
                subtotal_pagado += s.monto_pagado
                pendientes += 0 if s.pagado else 1
            cargos.append({
                'concepto': 'solvencia', 'concepto_nombre': 'Solvencia', 'nivel': 'alumno',
                'items': items, 'subtotal_usd': str(subtotal), 'subtotal_pagado_usd': str(subtotal_pagado),
                'saldo_usd': str(subtotal - subtotal_pagado), 'pendientes': pendientes,
            })
            deuda_total += subtotal - subtotal_pagado
            pagado_total += subtotal_pagado
            cargos_pendientes += pendientes

        # Cargos especiales (por tipo de cargo, el representante puede tener varios)
        cargos_especiales = list(
            CuotaProyectoInversion.objects.filter(representante=representante)
            .select_related('tipo_concepto').order_by('tipo_concepto__nombre', 'numero_cuota')
        )
        por_tipo = {}
        for ce in cargos_especiales:
            por_tipo.setdefault(ce.tipo_concepto_id, []).append(ce)
        for tipo_id, cuotas in por_tipo.items():
            tipo = cuotas[0].tipo_concepto
            items, subtotal, subtotal_pagado, pendientes = [], Decimal('0.00'), Decimal('0.00'), 0
            for ce in cuotas:
                saldo = ce.monto_usd - ce.monto_pagado
                estado = 'pagado' if ce.pagado else ('parcial' if ce.monto_pagado > 0 else 'pendiente')
                descripcion = tipo.nombre
                if tipo.numero_cuotas > 1:
                    descripcion = f"{tipo.nombre} — Cuota {ce.numero_cuota} de {tipo.numero_cuotas}"
                items.append({
                    'descripcion': descripcion,
                    'alumno': None,
                    'monto_usd': str(ce.monto_usd), 'monto_pagado_usd': str(ce.monto_pagado),
                    'saldo_usd': str(saldo), 'estado': estado, 'fecha_pago': ce.fecha_pago,
                })
                subtotal += ce.monto_usd
                subtotal_pagado += ce.monto_pagado
                pendientes += 0 if ce.pagado else 1
            cargos.append({
                'concepto': f'cargo_especial:{tipo_id}', 'concepto_nombre': tipo.nombre, 'nivel': 'representante',
                'items': items, 'subtotal_usd': str(subtotal), 'subtotal_pagado_usd': str(subtotal_pagado),
                'saldo_usd': str(subtotal - subtotal_pagado), 'pendientes': pendientes,
            })
            deuda_total += subtotal - subtotal_pagado
            pagado_total += subtotal_pagado
            cargos_pendientes += pendientes

        # Historial de pagos (paginado) — el OR con representante_documento
        # cubre los pagos retroactivos que guardan la cédula del
        # representante como texto suelto (no enlazados por FK alumno).
        historial_qs = Pago.objects.filter(
            Q(alumno__representante=representante) | Q(representante_documento=representante.cedula)
        ).distinct().select_related('alumno', 'banco_receptor', 'usuario_receptor').order_by('-fecha_pago')
        historial_qs = filtrar_por_sede(request.user, historial_qs, campo='sede')

        paginator = StandardResultsPagination()
        pagina = paginator.paginate_queryset(historial_qs, request, view=self)
        historial_response = paginator.get_paginated_response(PagoSerializer(pagina, many=True).data)

        return Response({
            'representante': {
                'id': representante.id,
                'nombre': f"{representante.nombre} {representante.apellido}",
                'cedula': representante.cedula,
                'telefono': representante.telefono,
                'correo': representante.correo,
            },
            'alumnos': alumnos_data,
            'cargos': cargos,
            'historial_pagos': historial_response.data,
            'totales': {
                'deuda_total_usd': str(deuda_total),
                'pagado_total_usd': str(pagado_total),
                'cargos_pendientes': cargos_pendientes,
            },
        })
